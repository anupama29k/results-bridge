"""Generate the M8 plate-scoring fixtures (deterministic; re-run to rebuild).

    python tests/fixtures/make_plate_fixtures.py

Writes Tecan Spark exports in the same layout as test_tecan_spark.xlsx (so
they go through the real parser), plus the plate maps they are read against.
Plate maps follow robotic-assays protocols/opentrons/plate_maps/ (cols 1-2
standards, cols 3-12 samples in duplicate) with QC / spike wells added.

Signal models are realistic, not idealised:
  BCA   -- curvilinear response (OD = 0.09 + 1.05e-3 x - 1.6e-7 x^2), 2% CV
  ELISA -- 4PL (a 0.05, b 1.2, c 0.10, d 2.8; x relative to top standard), 4% CV
"""
import json
import pathlib
import random

import openpyxl

HERE = pathlib.Path(__file__).parent
ROWS = "ABCDEFGH"


def write_spark(values: dict, path: pathlib.Path, wavelength: int):
    wb = openpyxl.load_workbook(HERE / "test_tecan_spark.xlsx")
    ws = wb.active
    ws["A3"] = "Protocol:"
    ws["B3"] = f"96-well Absorbance {wavelength}nm"
    ws["A4"] = f"Wavelength: {wavelength} nm"
    for r_i, r in enumerate(ROWS):
        for c in range(1, 13):
            ws.cell(row=6 + r_i, column=1 + c, value=round(values[f"{r}{c}"], 3))
    wb.save(path)


def layout(top, unit, dilution, n=40):
    pm = {"_meta": {"conc_unit": unit, "layout": "std cols 1-2, samples in duplicate cols 3-12"}}
    for c in (1, 2):
        for i, r in enumerate(ROWS):
            pm[f"{r}{c}"] = ({"role": "blank", "conc": 0} if r == "H" else
                             {"role": "standard", "conc": round(top / 2 ** i, 6), "group": f"STD{i + 1}"})
    for i in range(n):
        col = 3 + 2 * (i // 8)
        for c in (col, col + 1):
            pm[f"{ROWS[i % 8]}{c}"] = {"role": "sample", "group": f"S{i + 1:02d}", "dilution": dilution}
    return pm


def wells_of(pm, group):
    return [w for w, e in pm.items() if not w.startswith("_") and e.get("group") == group]


def bca():
    rng = random.Random(2)
    od = lambda x: 0.09 + 1.05e-3 * x - 1.6e-7 * x * x
    pm = layout(2000, "ug/mL", 1)
    true = {f"S{i:02d}": rng.uniform(100, 1500) for i in range(1, 41)}
    # S40 is a QC of known concentration
    for w in wells_of(pm, "S40"):
        pm[w] = {"role": "qc", "group": "QC1", "expected": round(true["S40"], 2), "dilution": 1}
    vals = {}
    for w, e in pm.items():
        if w.startswith("_"):
            continue
        x = e.get("conc", 0) if e["role"] in ("standard", "blank") else true[
            "S40" if e["role"] == "qc" else e["group"]]
        vals[w] = od(x) * rng.gauss(1, 0.02)
    return pm, vals


def elisa():
    rng = random.Random(5)
    a, b, c, d = 0.05, 1.2, 0.10, 2.8
    sig = lambda x: d + (a - d) / (1 + (x / c) ** b) if x > 0 else a
    pm = layout(1.0, "x top standard", 3)
    true = {f"S{i:02d}": rng.uniform(0.03, 1.2) for i in range(1, 41)}   # neat, relative
    true["S39"] = true["S38"] + 0.30                                       # S38 spiked with 0.30
    for w in wells_of(pm, "S39"):
        pm[w] = {"role": "positive_control", "group": "P01", "base_group": "S38",
                 "spike": 0.30, "dilution": 3}
    for w in wells_of(pm, "S40"):
        pm[w] = {"role": "qc", "group": "QC1", "expected": round(true["S40"], 4), "dilution": 3}
    vals = {}
    for w, e in pm.items():
        if w.startswith("_"):
            continue
        if e["role"] in ("standard", "blank"):
            x = e.get("conc", 0)
        else:
            src = {"qc": "S40", "positive_control": "S39"}.get(e["role"], e["group"])
            x = true[src] / 3
        vals[w] = max(sig(x) * rng.gauss(1, 0.04), 0.001)
    return pm, vals


if __name__ == "__main__":
    for name, (pm, vals), wl in (("bca", bca(), 562), ("elisa", elisa(), 450)):
        write_spark(vals, HERE / f"{name}_spark_plate.xlsx", wl)
        (HERE / f"{name}_plate_map.json").write_text(json.dumps(pm, indent=2) + "\n")
        print(f"wrote {name}_spark_plate.xlsx and {name}_plate_map.json")
